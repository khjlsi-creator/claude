# -*- coding: utf-8 -*-
"""레지스트리의 쪽 레이아웃(PG-)을 와이어프레임 PNG와 편집 가능한 PPTX(A4 세로)로 그린다."""
import os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
LAY = os.path.join(os.path.dirname(HERE), 'layouts')
PREV = os.path.join(LAY, 'preview')
PPTX_OUT = os.path.join(LAY, 'KORAD_경영평가_레이아웃_카탈로그.pptx')
PW, PH = 210.0, 297.0
FONT_PATH = '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'

# 블록 종류별 색 (원본 흑백 쪽을 KORAD 청색 계열 토큰으로 매핑)
FILL = dict(chapter_title='FFFFFF', eval_item_box='E6F4FF', bp_banner='BDE0FF', section_bar='001A7A',
            subhead='FFFFFF', labeled_block='EFF9FF', insight_bar='DEEEFF', arrow_caption='E1F4FF',
            rolling_block='E1F9FF', matrix_plus_channels='E6FCF9', tf_module='DEEEFF',
            divider_chevron='BDE0FF', ribbon_title='BDE0FF')
DARK = {'section_bar'}


def _blocks(reg, pid):
    pg = reg['pages'][pid]
    for b in pg['blocks']:
        lay = reg['blocks'][b['layout']]
        yield b, lay


def _abs(b):
    x0, x1 = b['x_mm']; y0, y1 = b['y_mm']
    return x0, y0, x1 - x0, y1 - y0


def render_png(reg, pid, scale=4.0):
    pg = reg['pages'][pid]
    W, H = int(PW * scale), int(PH * scale)
    im = Image.new('RGB', (W, H), 'white')
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(FONT_PATH, int(2.6 * scale))
    fs = ImageFont.truetype(FONT_PATH, int(2.0 * scale))
    d.rectangle([15 * scale, 10 * scale, 195 * scale, 287 * scale], outline='#BBBBBB')  # 여백 15/10mm
    for b, lay in _blocks(reg, pid):
        x, y, w, h = _abs(b)
        k = lay['kind']
        d.rectangle([x * scale, y * scale, (x + w) * scale, (y + h) * scale], fill='#' + FILL.get(k, 'FFFFFF'), outline='#0061DE')
        lw = lay.get('label_w') or 0
        if lw:
            lwm = lw * 210.0 / 1240
            d.rectangle([x * scale, y * scale, (x + lwm) * scale, (y + h) * scale], fill='#001A7A')
        for r in lay['regions']:
            rx, ry, rw, rh = r['rel_mm']
            d.rectangle([(x + rx) * scale, (y + ry) * scale, (x + rx + rw) * scale, (y + ry + rh) * scale], outline='#5D7280')
            d.text(((x + rx) * scale + 3, (y + ry) * scale + 2), r['name'], fill='#3B4951', font=fs)
        tc = 'white' if (k in DARK or lw) else '#000359'
        d.text((x * scale + 4, y * scale + 3), '%s %s %s' % (b['layout'], k, b.get('label') or ''), fill=tc, font=f)
    d.text((15 * scale, 2 * scale), '%s  %s  (%s)' % (pid, pg['title'], ','.join(pg['instances'])), fill='black', font=f)
    os.makedirs(PREV, exist_ok=True)
    out = os.path.join(PREV, '%s.png' % pid)
    im.save(out)
    return out


def _mm(v):
    return Emu(int(round(v * 36000)))


def _rect(s, x, y, w, h, fill, line='0061DE', text='', size=7, color='000359', bold=False, align=PP_ALIGN.LEFT, name=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, _mm(x), _mm(y), _mm(max(w, 0.5)), _mm(max(h, 0.5)))
    if name:
        sh.name = name
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Pt(0.5)
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = _mm(1); tf.margin_top = tf.margin_bottom = _mm(0.5)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = 'Pretendard'
    r.font.color.rgb = RGBColor.from_string(color)
    return sh


def render_pptx(reg):
    prs = Presentation()
    prs.slide_width, prs.slide_height = _mm(PW), _mm(PH)
    blank = prs.slide_layouts[6]
    # 1쪽: 현황표
    s = prs.slides.add_slide(blank)
    _rect(s, 15, 10, 180, 8, None, None, 'KORAD 경영평가 레이아웃 카탈로그 (누적)', 15, '000359', True, name='title')
    rows = ['프레임 %d종 · 블록 레이아웃 %d종 · 쪽 레이아웃 %d종' % (len(reg['frames']), len(reg['blocks']), len(reg['pages'])), '']
    for k, v in reg['pages'].items():
        rows.append('%s  %s  [%s]' % (k, v['title'], ','.join(v['instances'])))
    rows.append('')
    for k, v in reg['blocks'].items():
        rows.append('%s  %s  사용 %s' % (k, v['fp'], ','.join(v['used_in'])))
    _rect(s, 15, 20, 180, 265, None, None, '\n'.join(rows), 8, '3B4951', name='catalog')
    # 쪽 레이아웃별 와이어프레임 (도형 = 편집 가능)
    for pid, pg in reg['pages'].items():
        s = prs.slides.add_slide(blank)
        _rect(s, 15, 2, 180, 6, None, None, '%s  %s  (%s)' % (pid, pg['title'], ','.join(pg['instances'])), 10, '000359', True, name='page-title')
        _rect(s, 15, 10, 180, 277, None, 'BBBBBB', name='margin-15-10')
        for b, lay in _blocks(reg, pid):
            x, y, w, h = _abs(b)
            k = lay['kind']
            dark = k in DARK
            _rect(s, x, y, w, h, FILL.get(k, 'FFFFFF'), text='%s %s %s' % (b['layout'], k, b.get('label') or ''), size=7,
                  color='FFFFFF' if (dark or (lay.get('label_w') or 0)) else '000359', bold=True, name='%s %s' % (b['layout'], k))
            lw = (lay.get('label_w') or 0) * 210.0 / 1240
            if lw:
                _rect(s, x, y, lw, h, '001A7A', None, name='label-col')
            for r in lay['regions']:
                rx, ry, rw, rh = r['rel_mm']
                _rect(s, x + rx, y + ry, rw, rh, None, '5D7280', r['name'], 6, '3B4951', name=r['name'])
    prs.save(PPTX_OUT)
    return PPTX_OUT


def render_all(reg):
    for pid in reg['pages']:
        print('PNG', render_png(reg, pid))
    print('PPTX', render_pptx(reg))
