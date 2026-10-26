import sys, json, warnings
warnings.filterwarnings('ignore')
import pymupdf as fitz
from pptx import Presentation
pptx, pdf, out = sys.argv[1:4]
prs = Presentation(pptx)
doc = fitz.open(pdf)
meas = {}
for i, (slide, page) in enumerate(zip(prs.slides, doc), 1):
    tbls = sorted([sh for sh in slide.shapes if sh.has_table], key=lambda t: t.top)
    for ti, sh in enumerate(tbls):
        x0 = sh.left / 12700.0; x1 = (sh.left + sh.width) / 12700.0
        ytop = sh.top / 12700.0 - 1.5
        ybot = (tbls[ti + 1].top / 12700.0 - 1.5) if ti + 1 < len(tbls) else 9999
        ys = set()
        for d in page.get_drawings():
            col = d.get('color')
            if not col or not (abs(col[0]) < 0.03 and abs(col[1] - 0.557) < 0.04 and abs(col[2] - 1) < 0.03):
                continue
            for it in d['items']:
                if it[0] == 'l':
                    a, b = it[1], it[2]
                    if abs(a.y - b.y) < 0.3 and abs(a.x - b.x) > 20 and min(a.x, b.x) >= x0 - 3 and max(a.x, b.x) <= x1 + 3 and ytop <= a.y < ybot:
                        ys.add(round(a.y, 1))
        ys = sorted(ys)
        merged = []
        for y in ys:
            if not merged or y - merged[-1] > 1.5:
                merged.append(y)
        heights = [round((b - a) * 0.352778, 2) for a, b in zip(merged, merged[1:])]
        nrows = len(sh.table.rows)
        if len(heights) == nrows:
            meas[f'{i}|{sh.name}'] = heights
        else:
            print('MISMATCH', i, sh.name, len(heights), nrows)
json.dump(meas, open(out, 'w'), ensure_ascii=False)
print('measured', len(meas))
