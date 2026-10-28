# -*- coding: utf-8 -*-
"""디자인 시스템 미리보기용 HTML 조각을 만든다(그림 없음): 블록 견본 · RL 레이아웃 · PL 쪽 조합 · 레시피 카탈로그 쪽."""
import json, sys
from pptx import Presentation
import slots
from compose import compose
from layouts import LAYOUTS, PARTS, PAGES, K, PL_NODES
from recipes import NOTE
from layouts import catalog_spec
from pptx2html import slide_html, CSS

K2 = 960 / 1380.0
PAGE_H = 960 * 297 / 210


def region(slide, y_img, h_img, x0_img, w_img=1150, pad=6):
    y0 = (y_img - 48) * K2 - pad; y1 = (y_img - 48 + h_img) * K2 + pad; x0 = x0_img * K2 - pad
    w = w_img * K2 + 2 * pad
    return slide_html(slide, y0, y1, x0), round(w), round(y1 - y0)


def full(slide):
    return slide_html(slide), 960, round(PAGE_H)


SKIP = {'점', '값', '연도', '단위', '막대', '번호', '행 틀', '카드', '카드 머리', '박스 외곽', '그룹 틀', '분야 박스', '패널 외곽', '본문 틀', '단계 삼각형', '형광 밑줄'}


def generic_slots(slide, pl):
    """slots가 없는 신규 블록: 블록 영역 안에서 글이 들어가는 도형을 좌표 칸으로 돌려준다(블록 왼쪽 위 기준 격자 px)."""
    k = 1380 / 210 / 36000
    x0, y0, w, h = pl['x0'], pl['y'], 1150, pl['h']
    out = [('블록 전체', 0, 0, w, h)]
    seen = {}
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip() or sh.name in SKIP:
            continue
        px, py = sh.left * k, sh.top * k + 48
        if not (x0 - 2 <= px <= x0 + w and y0 - 2 <= py <= y0 + h):
            continue
        n = seen.get(sh.name, 0) + 1; seen[sh.name] = n
        if n > 4:
            continue
        out.append((f'{sh.name} {n}' if n > 1 or sum(1 for q in slide.shapes if q.name == sh.name) > 1 else sh.name, px - x0, py - y0, sh.width * k, sh.height * k))
    return out


def build(tmp):
    out = {'blocks': {}, 'rl': {}, 'pl': {}, 'recipe': [], 'css': CSS}
    def run(LS, tmp_name, record):
        nodes = []
        for i, L in enumerate(LS):
            if i: nodes.append(dict(t='newpage'))
            nodes += L['nodes']
        d, f = compose(dict(start_no=93, first='odd', note=NOTE, nodes=nodes), tmp + '/' + tmp_name)
        prs = Presentation(tmp + '/' + tmp_name)
        by_page = {}
        for pl, blk in zip(f.placed, f.placed_blocks):
            by_page.setdefault(pl['page'], []).append((pl, blk))
        rl = {}
        for i, L in enumerate(LS):
            items = by_page[i + 1]; sl = prs.slides[i]
            y0 = items[0][0]['y']; ymax = max(p['y'] + p['h'] for p, _ in items); x0 = items[0][0]['x0']
            if record:
                html, w, h = region(sl, y0, ymax - y0, x0)
                sls = []
                for pl, blk in items:
                    for name, x, y, ww, hh in (blk.slots(1150) if hasattr(blk, 'slots') else generic_slots(sl, pl)):
                        sls.append(dict(name=name, x=round((x * K2) + 6), y=round((pl['y'] - y0 + y) * K2 + 6), w=round(ww * K2), h=round(hh * K2), block=pl['cls'], mm=[round(x * 210 / 1380, 1), round((pl['y'] - 48 + y) * 210 / 1380, 1), round(ww * 210 / 1380, 1), round(hh * 210 / 1380, 1)]))
                out['rl'][L['id']] = dict(id=L['id'], name=L['name'], when=L['when'], src=L['src'], html=html, w=w, h=h, slots=sls,
                                          blocks=[dict(cls=p['cls'], h=round(p['h'])) for p, _ in items])
            rl[L['id']] = [(sl, p) for p, _ in items]
        return rl
    run(LAYOUTS, 'rl.pptx', True)
    rl = run(PARTS, 'parts.pptx', False)
    prs = Presentation(tmp + '/rl.pptx')
    def blk(rid, idx, cls=None):
        sl, p = rl[rid][idx]
        assert cls is None or p['cls'] == cls, (rid, idx, p['cls'])
        return region(sl, p['y'], p['h'], p['x0'])
    B = out['blocks']
    B['chapter-band'] = blk('L01', 0); B['eval-band-2line'] = blk('L02', 0); B['num-bar'] = blk('L03', 0); B['sub-head'] = blk('L04', 0)
    B['ladder-env'] = blk('L05', 0, 'Ladder'); B['ladder-side'] = blk('L06', 0, 'Ladder'); B['ladder-steps'] = blk('L07', 0, 'Ladder')
    B['ladder-banner-photo'] = blk('L09', 0, 'Ladder'); B['ladder-result'] = blk('L14', 0, 'Ladder'); B['ladder-photos'] = blk('L11', 0, 'Ladder')
    B['ladder-cause'] = blk('L12', 0, 'Ladder')
    B['label-rows-strategy'] = blk('L15', 0, 'LabelRows'); B['label-rows-check'] = blk('L16', 0, 'LabelRows'); B['label-rows-vision'] = blk('L17', 0, 'LabelRows')
    B['label-rows-nested'] = blk('L18', 0, 'LabelRows'); B['twin-panel'] = blk('L19', 0, 'TwinPanel'); B['three-panel'] = blk('L20', 0, 'ThreePanel')
    B['data-table'] = blk('L21', 0, 'DataTable'); B['result-1'] = blk('L22', 0, 'Result'); B['result-2'] = blk('L22', 1, 'Result'); B['result-prog'] = blk('L22', 2, 'Result')
    B['connector'] = blk('L23', 0, 'Connector'); B['note'] = blk('L16', 1, 'Note')
    # 사다리 + 연결 + 점검
    # PL
    for pid, letter, name, flow, src in PAGES:
        d2, f2 = compose(dict(start_no=93, first='odd', note=NOTE, nodes=PL_NODES[pid]), f'{tmp}/{pid.lower()}.pptx')
        pr = Presentation(f'{tmp}/{pid.lower()}.pptx')
        assert len(pr.slides) == 1, f'{pid}는 한 쪽이어야 합니다({len(pr.slides)}쪽): 대목을 줄이세요'
        fill = max(p['y'] + p['h'] for p in f2.placed) - 188
        assert fill / (1895 - 188) > 0.85, f'{pid}는 한 쪽을 85% 이상 채워야 합니다({fill / 17.07:.0f}%)'
        out_fill = fill / (1895 - 188)
        out['pl'][pid] = dict(id=pid, name=name, flow=flow, src=src, fill=round(out_fill * 100), pages=[full(s) for s in pr.slides],
                              blocks=[dict(page=p['page'], cls=p['cls']) for p in f2.placed])
        for sl, p in zip([pr.slides[q['page'] - 1] for q in f2.placed], f2.placed):
            if pid == 'PL05' and p['cls'] == 'EvalBand': B['eval-band-1line'] = region(sl, p['y'], p['h'], p['x0'])
    # 쪽 마스터: 머리·하단
    sl = prs.slides[0]
    odd = slide_html(sl, 0, 200, 0); even_prs = Presentation(tmp + '/pl02.pptx') if False else None
    B['master-odd'] = (slide_html(sl, 0, 110, 0) , 960, 110)
    B['master-odd-foot'] = (slide_html(sl, PAGE_H - 70, PAGE_H, 0), 960, 70)
    # 카탈로그
    d3, f3 = compose(catalog_spec(), tmp + '/cat.pptx')
    pr = Presentation(tmp + '/cat.pptx')
    out['recipe'] = [full(s) for s in pr.slides]
    even = [s for i, s in enumerate(pr.slides) if i % 2 == 1][0]
    B['master-even'] = (slide_html(even, 0, 110, 0), 960, 110); B['master-even-foot'] = (slide_html(even, PAGE_H - 70, PAGE_H, 0), 960, 70)
    return out


if __name__ == '__main__':
    import os
    tmp = sys.argv[1]; os.makedirs(tmp, exist_ok=True)
    o = build(tmp)
    json.dump(o, open(tmp + '/snips.json', 'w'), ensure_ascii=False)
    print(len(o['rl']), 'RL', len(o['pl']), 'PL', len(o['blocks']), 'blocks', len(o['recipe']), 'recipe pages')
